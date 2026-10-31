"""Stage 7: offline CV comparison and fail-closed staging checks."""
import json
import sys
from pathlib import Path
from docx import Document
import pytest

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from standalone.profile_cv_adapter import (
    _eligible, incompatible_numeric_context, ProfileEvidenceError,
    profile_evidence_mode, preserve_drive_core,
)
from standalone.profile_stage7_compare import compare_docx, main
from standalone.staging_profile_cv_handoff import ProfileStage3Sources


def make_docx(path, paragraphs, *, table=False, heading=False):
    doc=Document()
    if heading:
        doc.add_heading("Professional Experience",level=2)
    for p in paragraphs:
        doc.add_paragraph(p)
    if table:
        t=doc.add_table(rows=1,cols=2)
        t.cell(0,0).text="Skill"
        t.cell(0,1).text="Application"
    doc.save(path)
    return path


def test_numeric_context_disallows_revenue_claim_based_only_on_satisfaction():
    original="Improved client satisfaction by 40%. Reviewed 200+ companies."
    assert incompatible_numeric_context("Increased revenue by 40%.",original)
    assert not incompatible_numeric_context("Improved client satisfaction by 40%.",original)
    record={
        "source":{"kind":"master","fileId":"source-document-master-0123456"},
        "cvEligible":True,"status":"verified",
        "classificationReviewed":True,"metricsReviewed":True,
        "originalText":original,"metrics":["40%"],
        "cvBullet":"Increased revenue by 40%.",
        "star":{},"section":"experience",
    }
    assert not _eligible(record,"source-document-master-0123456")
    record["cvBullet"]="Improved client satisfaction by 40%."
    assert _eligible(record,"source-document-master-0123456")


def test_offline_cv_diff_shows_unsupported_new_numbers_and_structure(tmp_path):
    a=make_docx(tmp_path/"drive.docx",
                ["Client coordination and reporting. Reviewed 200+ companies."],heading=True)
    b=make_docx(tmp_path/"profile.docx",
                ["Client coordination and reporting. Reviewed 200+ companies.",
                 "Increased revenue by 40%."],table=True)
    result=compare_docx(a,b,"Company requires client reporting and investment analysis.")
    assert "40%" in result["numericClaimsAdded"]
    assert result["structureChanged"] is True
    assert result["requiresHumanReview"]
    assert result["jdKeywordCoverageProxy"]["jdTermsChecked"]>0
    assert result["driveStructure"]["tables"]==0
    assert result["profileStructure"]["tables"]==1
    assert result["notAutomaticallyVerified"]


def test_offline_cv_diff_identical_baseline_no_claim_or_format_change(tmp_path):
    a=make_docx(tmp_path/"a.docx",["Reviewed 200+ companies.","Produced financial analysis."])
    b=make_docx(tmp_path/"b.docx",["Reviewed 200+ companies.","Produced financial analysis."])
    result=compare_docx(a,b)
    assert not result["numericClaimsAdded"]
    assert not result["numericClaimsRemoved"]
    assert not result["structureChanged"]
    assert not result["requiresHumanReview"]


def test_comparison_cli_writes_local_report_and_requires_both_cover_letters(tmp_path):
    cv_a=make_docx(tmp_path/"drive.docx",["Reviewed 200+ companies."])
    cv_b=make_docx(tmp_path/"profile.docx",["Reviewed 200+ companies."])
    cl_a=make_docx(tmp_path/"drive-cl.docx",["I analysed market data."])
    cl_b=make_docx(tmp_path/"profile-cl.docx",["I analysed financial data."])
    report=tmp_path/"subfolder"/"report.json"
    assert main(["--drive-cv",str(cv_a),"--profile-cv",str(cv_b),
                 "--drive-cl",str(cl_a),"--profile-cl",str(cl_b),
                 "--report",str(report)])==0
    result=json.loads(report.read_text())
    assert result["coverLetter"] is not None
    assert result["humanSignoffRequired"] is True
    assert result["productionReleaseApproved"] is False
    assert report.is_file()
    with pytest.raises(SystemExit):
        main(["--drive-cv",str(cv_a),"--profile-cv",str(cv_b),
              "--drive-cl",str(cl_a),"--report",str(report)])


def test_stage7_production_mode_unavailable_without_external_staging_wrapper():
    assert profile_evidence_mode({})=="drive"
    for scope in ("production","",None):
        opts={"GARAMAAL_CV_EVIDENCE_MODE":"profile",
              "GARAMAAL_PROFILE_CV_CUTOVER_APPROVED":"true",
              "GARAMAAL_PROFILE_CV_STAGING_ISOLATED":"true"}
        if scope is not None:opts["GARAMAAL_PROFILE_DATA_SCOPE"]=scope
        with pytest.raises(ProfileEvidenceError,match="isolated staging"):
            profile_evidence_mode(opts)


def test_stage7_wrapper_uses_original_drive_sources_and_unchanged_runtime_by_default(tmp_path):
    selected=tmp_path/"Selected_Evidence.md"
    selected.write_text("original-drive-source")
    class Legacy:
        last_stage3_telemetry=[]
        def stage3(self,*args,**kwargs):
            return [selected]
    wrap=ProfileStage3Sources(Legacy(),environ={})
    result=wrap.stage3("pooja",tmp_path,None,None,"STAGE2_HANDOFF")
    assert result[0].read_text()=="original-drive-source"
    assert wrap.last_profile_evidence_audit is None


def test_stage7_staging_legacy_core_fails_closed_if_guardrails_missing():
    src="# SELECTED EVIDENCE\n## CORE-CHRONOLOGY\nReal chronology.\n## CORE-CREDENTIALS\nReal credentials.\n"
    with pytest.raises(ProfileEvidenceError,match="CORE-GUARDRAILS"):
        preserve_drive_core("profile evidence",src,"pooja")
