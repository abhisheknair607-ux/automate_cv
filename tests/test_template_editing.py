from pathlib import Path
from copy import deepcopy

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

from automate_cv.template_editing import edit_template, inspect_template


def fixture(path):
    doc = Document()
    doc.sections[0].left_margin = Cm(1.2)
    p = doc.add_paragraph('POOJA TIRUPATI')
    p.runs[0].bold = True
    p.runs[0].font.size = Pt(14)
    contact = doc.add_paragraph()
    link = OxmlElement('w:hyperlink')
    link.set(qn('w:anchor'),'LinkedIn')
    run = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text='LinkedIn'
    run.append(text); link.append(run); contact._p.append(link)
    p = doc.add_paragraph()
    r = p.add_run('Role\t2025–2026'); r.bold=True; r.font.size=Pt(10)
    p.paragraph_format.tab_stops.add_tab_stop(Cm(18.5))
    p = doc.add_paragraph('Supported verified work.')
    p.runs[0].bold=False; p.runs[0].font.size=Pt(10)
    p.paragraph_format.left_indent=Cm(.5)
    p.paragraph_format.first_line_indent=Cm(-.5)
    doc.save(path)


def test_template_geometry_hyperlink_tabs_and_indents_are_preserved(tmp_path):
    source, output=tmp_path/'source.docx',tmp_path/'new_cv.docx'
    fixture(source); original=source.read_bytes()
    before=Document(source)
    edit_template(source,output,[{'source':0},{'source':1},
        {'source':2,'runs':[{'text':'Updated role\t2025–2026'}]},
        {'source':3,'runs':[{'text':'Supported accurate revised wording.'}]}])
    after=Document(output)
    assert source.read_bytes()==original
    assert before.sections[0]._sectPr.xml==after.sections[0]._sectPr.xml
    assert before.paragraphs[1]._p.xml==after.paragraphs[1]._p.xml
    for i in (2,3): assert before.paragraphs[i]._p.pPr.xml==after.paragraphs[i]._p.pPr.xml
    assert after.paragraphs[2].text=='Updated role\t2025–2026'
    assert after.paragraphs[2].runs[0].bold is True
    assert after.paragraphs[3].runs[0].bold is False
    assert after.paragraphs[3].runs[0].font.size.pt==10
    assert after.paragraphs[3].runs[0].font.name=='Calibri'


def test_cloned_bullets_reuse_properties_without_blank_spacing(tmp_path):
    source, output=tmp_path/'source.docx',tmp_path/'new.docx'
    fixture(source)
    edit_template(source,output,[{'source':3,'runs':[{'text':'First verified bullet.'}]},
        {'source':3,'runs':[{'text':'Second verified bullet.'}]}])
    doc=Document(output)
    assert len(doc.paragraphs)==2
    assert doc.paragraphs[0]._p.pPr.xml==doc.paragraphs[1]._p.pPr.xml


def test_output_cannot_overwrite_source_or_use_invalid_prototype(tmp_path):
    source=tmp_path/'source.docx';fixture(source)
    with pytest.raises(ValueError,match='new file'):edit_template(source,source,[])
    with pytest.raises(ValueError,match='paragraph index'):
        edit_template(source,tmp_path/'new.docx',[{'source':99}])


def test_inspection_returns_stable_prototype_indices(tmp_path):
    source=tmp_path/'source.docx';fixture(source)
    rows=inspect_template(source)
    assert rows[0]['index']==0 and rows[0]['runs'][0]['size_pt']==14


def test_explicit_regular_text_overrides_bold_prototype(tmp_path):
    source, output=tmp_path/'source.docx',tmp_path/'new.docx'
    fixture(source)
    edit_template(source,output,[{'source':2,'runs':[{'text':'Regular detail','bold':False}]}])
    assert Document(output).paragraphs[0].runs[0].bold is False
