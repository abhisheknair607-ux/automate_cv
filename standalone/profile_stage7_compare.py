"""Read-only local CV / CL side-by-side inspection for Stage 7 staging.

Does not create or modify a Word document, send CV content to any API, call
the paid worker, or assert an actual ATS score/page count. Only compares
already-produced, locally supplied artifacts. Report contains no full CV text.
"""
from __future__ import annotations
import argparse
import json
import re
from collections import Counter
from pathlib import Path

from docx import Document
from standalone.profile_cv_adapter import NUMBER

STOP_WORDS = set("""about above after again against all also among and any are
as at be because been before below between both but by can could do does did
each else for from further had has have having how if in into is it its itself
just like may more most must not of on or our other out over own per should
so some such than that the their them then there these they this those through
to too under until up used using very was were what when where which while
who why will with within would you your role job experience required ability
skills responsibilities candidate essential preferred minimum desirable""".split())
WORD = re.compile(r"[a-z][a-z0-9+.-]{2,}", re.I)


def _terms(text: str) -> set[str]:
    return {w.lower() for w in WORD.findall(text) if w.lower() not in STOP_WORDS}


def _numbers(text: str) -> list[str]:
    return sorted({re.sub(r"[\s,]+", "", m.group(0)).lower()
                   for m in NUMBER.finditer(text)})


def _structure(doc: Document) -> dict:
    sections = []
    for sec in doc.sections:
        sections.append({
            "widthInches": round(sec.page_width.inches, 3),
            "heightInches": round(sec.page_height.inches, 3),
            "topMarginInches": round(sec.top_margin.inches, 3) if sec.top_margin else None,
            "bottomMarginInches": round(sec.bottom_margin.inches, 3) if sec.bottom_margin else None,
        })
    return {
        "sectionGeometry": sections,
        "tables": len(doc.tables),
        "paragraphs": len(doc.paragraphs),
        "headingStyles": sum(1 for p in doc.paragraphs
                             if p.style and p.style.name.lower().startswith("heading")),
        "explicitPageBreaks": len(doc.element.xpath('.//w:br[@w:type="page"]')),
        "sectionCount": len(doc.sections),
    }


def _text(doc: Document) -> str:
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.append(cell.text)
    for section in doc.sections:
        texts.extend(p.text for p in section.header.paragraphs)
        texts.extend(p.text for p in section.footer.paragraphs)
    return "\n".join(texts)


def compare_docx(drive_file: str | Path, profile_file: str | Path,
                 jd: str = "") -> dict:
    old = Document(str(drive_file))
    new = Document(str(profile_file))
    a, b = _text(old), _text(new)
    old_n, new_n = set(_numbers(a)), set(_numbers(b))
    old_terms, new_terms = _terms(a), _terms(b)
    jd_terms = _terms(jd) if jd else set()
    a_struct, b_struct = _structure(old), _structure(new)
    delta = {
        "textCharacterChange": len(b)-len(a),
        "numericClaimsAdded": sorted(new_n-old_n),
        "numericClaimsRemoved": sorted(old_n-new_n),
        "addedVocabularyCount": len(new_terms-old_terms),
        "removedVocabularyCount": len(old_terms-new_terms),
        "structureChanged": a_struct != b_struct,
        "driveStructure": a_struct, "profileStructure": b_struct,
        "possibleFormattingChange": a_struct != b_struct,
        "notAutomaticallyVerified": [
            "Numbers can be reused for different outcomes.",
            "A Word file does not establish rendered page count.",
            "Keyword coverage is not an ATS score.",
            "A human must verify employer, dates, chronology, ownership and source citations.",
        ],
    }
    if jd_terms:
        delta["jdKeywordCoverageProxy"] = {
            "jdTermsChecked": len(jd_terms),
            "driveMatched": len(jd_terms & old_terms),
            "profileMatched": len(jd_terms & new_terms),
            "newlyCovered": sorted(jd_terms & (new_terms-old_terms))[:35],
            "nowMissing": sorted(jd_terms & (old_terms-new_terms))[:35],
        }
    delta["requiresHumanReview"] = bool(
        delta["numericClaimsAdded"] or delta["numericClaimsRemoved"] or
        delta["possibleFormattingChange"]
    )
    return delta


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description="Compare existing local CV/CL DOCX outputs without modification")
    parser.add_argument("--drive-cv", type=Path, required=True)
    parser.add_argument("--profile-cv", type=Path, required=True)
    parser.add_argument("--drive-cl", type=Path)
    parser.add_argument("--profile-cl", type=Path)
    parser.add_argument("--jd-file", type=Path)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args(argv)
    if bool(args.drive_cl) != bool(args.profile_cl):
        parser.error("Provide both Drive and Profile cover letters or neither")
    for file in [args.drive_cv,args.profile_cv,args.drive_cl,args.profile_cl]:
        if file and (not file.is_file() or file.suffix.lower() != ".docx"):
            parser.error("All provided Word files must be existing DOCX files")
    jd = args.jd_file.read_text(encoding="utf-8") if args.jd_file else ""
    if len(jd)>200000:
        parser.error("JD is too large")
    result={"schema":"garamaal-staging-comparison-v1",
            "cv":compare_docx(args.drive_cv,args.profile_cv,jd),
            "coverLetter":compare_docx(args.drive_cl,args.profile_cl,jd)
            if args.drive_cl else None,
            "humanSignoffRequired":True,
            "productionReleaseApproved":False}
    args.report.parent.mkdir(parents=True,exist_ok=True)
    args.report.write_text(json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print("Local comparison report written. No document or remote system changed.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
