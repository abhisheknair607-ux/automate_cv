"""On-demand scoring for Fresh Maal; CV automation columns remain untouched."""
import argparse
import io
import json
import os
from decimal import Decimal, ROUND_HALF_UP

from docx import Document
from googleapiclient.discovery import build
from openai import OpenAI
from .config import settings
from .google_auth import credentials

CANDIDATES = {
    "abhishek": {
        "name": "Abhishek Nair",
        "summary": "SUMMARY_DOC_FILE_ID",
        "goal": "1KqnWYEi9GejUhkqiRnQ1D1G6DlAtnH_M",
        "score_col": "AL", "score_index": 37, "selected_index": 41,
    },
    "pooja": {
        "name": "Pooja Krishna Tirupati",
        "summary": "POOJA_SUMMARY_DOC_FILE_ID",
        "goal": "1Fyy2txzfuWHFjARgZjOETDplYHqGHy2o",
        "score_col": "AN", "score_index": 39, "selected_index": 42,
    },
}

INSTRUCTIONS = """
Rank jobs using the CURRENT candidate Ultimate Career Goal and their VERIFIED
full Summary Doc. Raw JD controls vacancy facts, Summary Doc candidate facts,
and current Career Goal controls strategic alignment. Never invent evidence.
For Pooja, CURRENT PRIMARY priority is front-office PRIVATE WEALTH, regulated
financial planning, investments and HNW advisory. Disregard superseded
PE/VC-primary strategy from older job ranking prompt.
For Abhishek, priority is Deals, Valuation, FDD, Corporate Finance, Private
Credit, PE and transaction-led investing.

Profile Fit /100: professional duties 30; mandatory qualifications and seniority
20; analytical, financial and software skills 15; sector, transactions and
product match 15; client/management/stakeholder skills 10; strength and
credibility of evidence 10.
Missing mandatory professional or regulated qualification normally caps
Profile Fit at 20. Missing essential sector experience caps it at 55 unless
transferability is explicitly permitted. Material seniority/direct experience
mismatch normally caps at 70. Projects never erase professional-experience
hard gates.
Career Alignment /100: judge whether the ACTUAL work and credible next move
advance the specific career-goal file, giving high credit to relevant live
experience, meaningful client/advisory/investment/transaction responsibility
and low credit to routine processing or pure compliance.
Consider all relevant source facts critically. Do not invent visa permissions.
Do not research employers. Return ONLY a JSON object with property "jobs":
an array of objects, exactly {"row": integer, "profile_fit": integer,
"career_alignment": integer}. Every input row exactly once. Integers 0–100.
No prose or markdown.
"""


def summary_doc(raw):
    doc = Document(io.BytesIO(raw))
    parts = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            text = " | ".join(cell.text.replace("\n", "; ").strip() for cell in row.cells)
            if text.strip(" |"):
                parts.append(text)
    return "\n".join(parts)


def parse_scores(output, expected):
    raw = str(output or "").strip()
    fence = chr(96) * 3
    if raw.startswith(fence):
        raw = raw.split("\n", 1)[1].rsplit(fence, 1)[0].strip()
    jobs = json.loads(raw)["jobs"]
    if not isinstance(jobs, list) or len(jobs) != len(expected):
        raise ValueError("Ranking response has unexpected number of rows.")
    scores = {}
    for item in jobs:
        row, fit, goal = item["row"], item["profile_fit"], item["career_alignment"]
        if type(row) is not int or row not in expected or row in scores:
            raise ValueError("Wrong or duplicate row in ranking result.")
        if type(fit) is not int or type(goal) is not int:
            raise ValueError("Ranking components must be integers.")
        if not 0 <= fit <= 100 or not 0 <= goal <= 100:
            raise ValueError("Ranking component outside 0–100.")
        combined = Decimal(fit) * Decimal(".55") + Decimal(goal) * Decimal(".45")
        scores[row] = int(combined.quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    if scores.keys() != expected:
        raise ValueError("Missing scored job.")
    return scores


def run_person(person, args, sheets, drive, client):
    config = CANDIDATES[person]
    source_id = os.getenv(config["summary"], "").strip()
    if not source_id:
        raise RuntimeError(f"{config['summary']} not configured; no scores written.")
    profile = summary_doc(drive.files().get_media(fileId=source_id).execute())
    goal = drive.files().get_media(fileId=config["goal"]).execute().decode("utf-8-sig")
    if len(profile) < 400 or len(goal) < 400:
        raise RuntimeError("Full candidate ranking source unavailable.")
    if person == "pooja" and "PRIVATE WEALTH" not in goal.upper():
        raise RuntimeError("Pooja's current Private Wealth goal cannot be verified.")
    sid, sheet = settings.freshmal_spreadsheet_id, settings.freshmal_sheet_name
    if not sid:
        raise RuntimeError("FRESHMAL_SPREADSHEET_ID not configured.")
    data = sheets.spreadsheets().values().get(
        spreadsheetId=sid, range=f"'{sheet}'!A3:AQ995",
        valueRenderOption="UNFORMATTED_VALUE",
    ).execute().get("values", [])
    pending = []
    for num, cells in enumerate(data, 3):
        def at(i):
            return cells[i] if len(cells) > i else ""
        company, role, url, jd = [at(i) for i in range(4)]
        if not (str(company).strip() and str(role).strip() and str(jd).strip()):
            continue
        selected = at(config["selected_index"])
        if args.selected_only and selected is not True and str(selected).upper() != "TRUE":
            continue
        prior = at(config["score_index"])
        if not args.refresh and prior not in ("", None):
            continue
        pending.append({
            "row": num, "company": str(company), "role": str(role),
            "url": str(url), "jd": str(jd),
        })
        if len(pending) >= args.max_jobs:
            break
    if not pending:
        print(f"{person}: no eligible jobs; no AI call.", flush=True)
        return 0

    saved = 0
    for start in range(0, len(pending), args.batch_size):
        group = pending[start:start + args.batch_size]
        prompt = json.dumps({
            "candidate": config["name"],
            "verified_summary_doc": profile,
            "current_career_goal": goal, "jobs": group,
        }, ensure_ascii=False)
        response = client.responses.create(
            model=os.getenv("RANK_MODEL", "gpt-5.6-terra"),
            instructions=INSTRUCTIONS, input=prompt,
            reasoning={"effort": "medium"}, max_output_tokens=1800,
        )
        scores = parse_scores(response.output_text, {x["row"] for x in group})
        updates = []
        for job in group:
            n = job["row"]
            live_rows = sheets.spreadsheets().values().get(
                spreadsheetId=sid, range=f"'{sheet}'!A{n}:D{n}",
                valueRenderOption="UNFORMATTED_VALUE",
            ).execute().get("values", [[]])
            live = live_rows[0] if live_rows else []
            if len(live) < 4 or tuple(map(str, live[:4])) != (
                job["company"], job["role"], job["url"], job["jd"]
            ):
                print(f"Row {n} edited during ranking; skipped.", flush=True)
                continue
            updates.append({
                "range": f"'{sheet}'!{config['score_col']}{n}",
                "values": [[scores[n]]],
            })
        if updates:
            sheets.spreadsheets().values().batchUpdate(
                spreadsheetId=sid,
                body={"valueInputOption": "RAW", "data": updates},
            ).execute()
            saved += len(updates)
    print(f"{person}: wrote {saved} scores; rank formula recalculates.", flush=True)
    return saved


def main():
    p = argparse.ArgumentParser(description="Manually rank jobs without CV generation")
    p.add_argument("--candidate", choices=["abhishek", "pooja", "both"], default="both")
    p.add_argument("--max-jobs", type=int, default=10)
    p.add_argument("--batch-size", type=int, default=3)
    p.add_argument("--selected-only", action="store_true")
    p.add_argument("--refresh", action="store_true")
    args = p.parse_args()
    if not 1 <= args.max_jobs <= 50 or not 1 <= args.batch_size <= 5:
        p.error("Use 1–50 jobs per candidate and batch size 1–5.")
    if not settings.openai_api_key:
        raise RuntimeError("OPENAI_API_KEY must be configured.")
    auth = credentials(settings)
    sheets = build("sheets", "v4", credentials=auth, cache_discovery=False)
    drive = build("drive", "v3", credentials=auth, cache_discovery=False)
    client = OpenAI(api_key=settings.openai_api_key)
    people = CANDIDATES if args.candidate == "both" else [args.candidate]
    count = sum(run_person(person, args, sheets, drive, client) for person in people)
    print(f"Done. {count} jobs ranked in main Fresh Maal. No CV workflow run.", flush=True)


if __name__ == "__main__":
    main()
