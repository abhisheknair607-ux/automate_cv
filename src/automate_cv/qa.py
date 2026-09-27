import re
import shutil
import subprocess
from pathlib import Path
from docx import Document
from pypdf import PdfReader


class QAError(RuntimeError):
    pass


def _artifact_score(path, kind):
    name = Path(path).name.lower()
    size = Path(path).stat().st_size if Path(path).exists() else 0
    if kind == 'cv':
        score = 0
        if 'cv' in name or 'resume' in name or 'résumé' in name:
            score += 20
        if 'cover' in name or 'letter' in name:
            score -= 100
    else:
        score = 0
        if 'cover' in name:
            score += 30
        if 'letter' in name:
            score += 20
        if 'cv' in name or 'resume' in name:
            score -= 50
    return score, size, name


def normalize(artifacts, cv_name, cl_name, output_dir):
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    docs = [Path(p) for p in artifacts if str(p).lower().endswith('.docx') and Path(p).exists()]
    if not docs:
        raise QAError('Stage 3 produced no DOCX artifacts.')

    cv_candidates = [p for p in docs if 'cover' not in p.name.lower() and 'letter' not in p.name.lower()]
    if not cv_candidates:
        raise QAError('Stage 3 produced no identifiable CV DOCX.')
    cv_source = sorted(cv_candidates, key=lambda p: _artifact_score(p, 'cv'), reverse=True)[0]
    cv_target = output_dir / cv_name
    shutil.copy2(cv_source, cv_target)

    cl_target = None
    if cl_name:
        cl_candidates = [p for p in docs if p != cv_source and ('cover' in p.name.lower() or 'letter' in p.name.lower())]
        if not cl_candidates:
            raise QAError('Cover letter requested but no identifiable cover-letter DOCX was produced.')
        cl_source = sorted(cl_candidates, key=lambda p: _artifact_score(p, 'cl'), reverse=True)[0]
        cl_target = output_dir / cl_name
        shutil.copy2(cl_source, cl_target)
    return cv_target, cl_target


def one_page(docx, render_dir):
    render_dir = Path(render_dir)
    render_dir.mkdir(parents=True, exist_ok=True)
    r = subprocess.run(
        ['libreoffice', '--headless', '--convert-to', 'pdf', '--outdir', str(render_dir), str(docx)],
        capture_output=True,
        text=True,
        timeout=120,
    )
    if r.returncode:
        raise QAError(r.stderr or 'LibreOffice rendering failed.')
    pdf = render_dir / (Path(docx).stem + '.pdf')
    if not pdf.exists():
        raise QAError('Rendered PDF missing.')
    pages = len(PdfReader(str(pdf)).pages)
    if pages != 1:
        raise QAError(f'{Path(docx).name} rendered to {pages} pages; expected 1.')


def _docx_text(path):
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            parts.append(' | '.join(cell.text for cell in row.cells))
    return '\n'.join(parts)


def _norm(text):
    text = str(text).replace('–', '-').replace('—', '-').replace('−', '-')
    text = re.sub(r'\s+', ' ', text.lower()).strip()
    return text


def _percent_values(text):
    return {m.group(1).lstrip('0') or '0' for m in re.finditer(r'\b(\d+(?:\.\d+)?)\s*%', text)}


def _money_values(text):
    values = set()
    pattern = re.compile(
        r'(?P<cur>[€$£])\s*(?P<num>\d[\d,]*(?:\.\d+)?)\s*(?P<mag>k|m|mn|million|bn|billion)?',
        re.I,
    )
    mult = {'': 1, 'k': 1_000, 'm': 1_000_000, 'mn': 1_000_000, 'million': 1_000_000, 'bn': 1_000_000_000, 'billion': 1_000_000_000}
    for m in pattern.finditer(text):
        try:
            number = float(m.group('num').replace(',', '')) * mult[(m.group('mag') or '').lower()]
            values.add((m.group('cur'), round(number, 2)))
        except Exception:
            continue
    return values


def _qualification_issues(cv_text, source_text):
    issues = []
    qualifications = ('ACCA', 'ACA', 'CIMA', 'CFA', 'CPA', 'QFA', 'MBA', 'MSc', 'BSc', 'BCom', 'PhD')
    statuses = ('qualified', 'member', 'candidate', 'affiliate', 'passed')
    cv_lower = cv_text.lower()
    src_lower = source_text.lower()
    for q in qualifications:
        if re.search(rf'\b{re.escape(q)}\b', cv_text, re.I) and not re.search(rf'\b{re.escape(q)}\b', source_text, re.I):
            issues.append(f'Qualification {q} appears in the CV but not in verified source material.')
        for status in statuses:
            q_then_status = re.search(rf'\b{re.escape(q)}\b[^\n.]{{0,80}}\b{status}\b', cv_text, re.I)
            status_then_q = re.search(rf'\b{status}\b[^\n.]{{0,80}}\b{re.escape(q)}\b', cv_text, re.I)
            if q_then_status or status_then_q:
                if q.lower() not in src_lower or status not in src_lower:
                    issues.append(f'{q} status term "{status}" is not supported by the verified source material.')
    return list(dict.fromkeys(issues))


def _tool_issues(cv_text, source_text):
    issues = []
    tools = ('SAP', 'BPC', 'Power BI', 'Tableau', 'Python', 'SQL', 'Alteryx', 'VBA', 'Excel', 'Oracle', 'SAS', 'Power Query', 'Power Pivot', 'Bloomberg')
    for tool in tools:
        if re.search(rf'(?<!\w){re.escape(tool)}(?!\w)', cv_text, re.I) and not re.search(rf'(?<!\w){re.escape(tool)}(?!\w)', source_text, re.I):
            issues.append(f'Tool/software {tool} appears in the CV but not in verified source material.')
    return issues


def factual_report(cv_path, source_text, report_path):
    """Deterministic high-risk fact check.

    This is intentionally conservative: it blocks unsupported currency amounts,
    percentages, qualification/status terms and named tools. Other numeric claims
    are surfaced as warnings for audit rather than automatically rejected.
    """
    cv_text = _docx_text(cv_path)
    source_text = source_text or ''
    critical = []
    warnings = []

    source_pcts = _percent_values(source_text)
    for value in sorted(_percent_values(cv_text)):
        if value not in source_pcts:
            critical.append(f'{value}% appears in the CV but not in verified source material.')

    source_money = _money_values(source_text)
    for cur, value in sorted(_money_values(cv_text)):
        if (cur, value) not in source_money:
            critical.append(f'{cur}{value:,.2f} equivalent appears in the CV but not in verified source material.')

    critical.extend(_qualification_issues(cv_text, source_text))
    critical.extend(_tool_issues(cv_text, source_text))

    metric_pattern = re.compile(r'\b\d+(?:\.\d+)?\+?\s+(?:years?|clients?|countries?|projects?|entities?|jurisdictions?|reports?|models?)\b', re.I)
    source_norm = _norm(source_text)
    for m in metric_pattern.finditer(cv_text):
        claim = _norm(m.group(0))
        if claim not in source_norm:
            warnings.append(f'Numeric claim requires review: {m.group(0)}')

    critical = list(dict.fromkeys(critical))
    warnings = list(dict.fromkeys(warnings))
    report_path = Path(report_path)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    lines = [
        '# Factual QA',
        '',
        f'- Critical unsupported high-risk claims: {len(critical)}',
        f'- Numeric review warnings: {len(warnings)}',
        '',
        '## Critical',
    ]
    lines.extend([f'- {x}' for x in critical] or ['- None'])
    lines.extend(['', '## Warnings'])
    lines.extend([f'- {x}' for x in warnings] or ['- None'])
    report_path.write_text('\n'.join(lines) + '\n', encoding='utf-8')
    return {'critical': critical, 'warnings': warnings, 'report_path': str(report_path)}
