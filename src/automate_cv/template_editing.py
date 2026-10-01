"""Reusable DOCX editing and rendering utilities for the Stage 3 container.

Preserve template geometry, paragraph/run formatting and package relationships.
The caller supplies verified text and uses the Formatting Master for exceptions.
"""
from copy import deepcopy
from pathlib import Path
import shutil
import subprocess

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def _text_run(prototype, item):
    element = OxmlElement('w:r')
    if prototype is not None and prototype.rPr is not None:
        element.append(deepcopy(prototype.rPr))
    props = element.get_or_add_rPr()
    if 'bold' in item:
        for old in list(props.findall(qn('w:b'))):
            props.remove(old)
        bold = OxmlElement('w:b')
        bold.set(qn('w:val'), '1' if item['bold'] else '0')
        props.append(bold)
    # Explicit Calibri prevents theme inheritance from changing newly inserted text.
    fonts = props.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        props.append(fonts)
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:' + k), 'Calibri')
    for k in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + k), None)
    # Text is literal. Tabs/newlines become real OOXML controls, never manual spacing.
    import re
    for value in re.split(r'(\t|\n)', str(item.get('text', ''))):
        if value == '\t':
            element.append(OxmlElement('w:tab'))
        elif value == '\n':
            element.append(OxmlElement('w:br'))
        elif value:
            node = OxmlElement('w:t')
            node.set(qn('xml:space'), 'preserve')
            node.text = value
            element.append(node)
    return element


def edit_template(template_path, output_path, blocks):
    """Write a duplicate using ordered template paragraph prototypes.

blocks: [{source: paragraph_index, runs: [{text, source_run?, bold?}]}].
Omit runs to retain the source paragraph intact (especially contact hyperlinks).
Repeat a source index to clone another bullet/role. Omit lower-ranked paragraphs
to remove them. runs inherit the indexed prototype's run properties; choose the
matching bold/regular prototype for each fragment. Paragraph formatting, borders,
numbering, tabs, page setup, headers/footers and package relationships are retained.
Never use this to manufacture evidence or an unapproved section/layout.
"""
    template_path, output_path = Path(template_path), Path(output_path)
    if template_path.resolve() == output_path.resolve():
        raise ValueError('Output must be a new file, never the source template.')
    doc = Document(template_path)
    if doc.tables:
        raise ValueError('Template uses tables; preserve those with a source-specific edit.')
    original = list(doc.paragraphs)
    prepared = []
    for block in blocks:
        index = int(block['source'])
        if index < 0 or index >= len(original):
            raise ValueError(f'Invalid template paragraph index: {index}')
        proto = original[index]
        element = deepcopy(proto._p)
        if 'runs' in block:
            for child in list(element):
                if child.tag != qn('w:pPr'):
                    element.remove(child)
            for item in block['runs']:
                run_index = int(item.get('source_run', 0))
                if run_index < 0 or (proto.runs and run_index >= len(proto.runs)):
                    raise ValueError(f'Invalid run index {run_index} for paragraph {index}')
                prototype = proto.runs[run_index]._r if proto.runs else None
                element.append(_text_run(prototype, item))
        prepared.append(element)
    body = doc.element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    section = body.find(qn('w:sectPr'))
    for element in prepared:
        if section is None:
            body.append(element)
        else:
            section.addprevious(element)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)
    return output_path


def inspect_template(template_path):
    """Return compact paragraph/run prototypes once, rather than dumping all XML."""
    doc = Document(template_path)
    return [{'index': i, 'text': p.text,
             'runs': [{'index': j, 'text': r.text, 'bold': r.bold,
                       'size_pt': r.font.size.pt if r.font.size else None}
                      for j, r in enumerate(p.runs)]}
            for i, p in enumerate(doc.paragraphs) if p.text.strip()]


def render_document(docx_path, output_dir):
    """Render PDF and PNGs for actual visual inspection; never claim visual QA."""
    docx_path, output_dir = Path(docx_path), Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    executable = shutil.which('libreoffice') or shutil.which('soffice')
    if not executable:
        raise RuntimeError('A DOCX renderer is unavailable; visual QA is not complete.')
    pdf = output_dir / (docx_path.stem + '.pdf')
    pdf.unlink(missing_ok=True)
    result = subprocess.run([executable, '--headless', '--convert-to', 'pdf',
                             '--outdir', str(output_dir), str(docx_path)],
                            capture_output=True, text=True, timeout=120)
    if result.returncode or not pdf.is_file():
        raise RuntimeError(result.stderr or 'DOCX rendering failed.')
    pngs = []
    try:
        import fitz
        with fitz.open(pdf) as rendered:
            count = len(rendered)
            for i, page in enumerate(rendered):
                path = output_dir / f'{docx_path.stem}-page-{i+1}.png'
                page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(path)
                pngs.append(str(path))
    except ImportError:
        from pypdf import PdfReader
        count = len(PdfReader(str(pdf)).pages)
        converter = shutil.which('pdftoppm')
        if not converter:
            raise RuntimeError('PDF produced but PNG renderer unavailable; inspect the PDF.')
        prefix = output_dir / (docx_path.stem + '-page')
        subprocess.run([converter, '-png', '-r', '110', str(pdf), str(prefix)], check=True)
        pngs = [str(p) for p in sorted(output_dir.glob(prefix.name + '-*.png'))]
    return {'pdf': str(pdf), 'page_count': count, 'page_images': pngs,
            'visual_inspection_required': True}
