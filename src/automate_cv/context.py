from pathlib import Path
from docx import Document


def cv_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs if p.text]
    for table in d.tables:
        for row in table.rows:
            parts.append(' | '.join(c.text for c in row.cells))
    return '\n'.join(parts)


def write_context(path, app, stage1, stage2, stage3, cv_path, cv_link):
    text = f'''# Application Context — {app.candidate_name} — {app.company} — {app.designation}

Candidate: {app.candidate_name}
Candidate key: {app.candidate_key}
Application ID: {app.application_id}
Company: {app.company}
Designation: {app.designation}
Location: {app.location or 'Not supplied'}
Vacancy: {app.application_link}
Current CV: {cv_link}
Make Cover Letter: {str(app.make_cl).upper()}
Web Search: {str(app.web_search).upper()}

## Continuation instruction
Continue only for the candidate named above, from the current CV and verified application context. Do not mix evidence from another candidate. Do not restart Stages 1–3 unless explicitly requested. Do not invent experience, metrics, qualifications, responsibilities, ownership, tools or achievements.

## Raw Job Description
{app.raw_jd}

## Stage 1
{stage1}

## Stage 2
{stage2}

## Stage 3 / QA notes
{stage3}

## Current CV Text
{cv_text(cv_path)}
'''
    path = Path(path)
    path.write_text(text, encoding='utf-8')
    return path
